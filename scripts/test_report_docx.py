import os
import sys
import unittest
import docx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

class TestReportDocx(unittest.TestCase):
    def setUp(self):
        self.repo_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        self.doc_path = os.path.join(self.repo_dir, "BAO_CAO_TTTN_Smart_Desk_Assistant.docx")

    def test_report_file_exists(self):
        self.assertTrue(os.path.exists(self.doc_path), "BAO_CAO_TTTN_Smart_Desk_Assistant.docx must exist")
        self.assertGreater(os.path.getsize(self.doc_path), 30000, "Report file size must be > 30KB")

    def test_report_content_and_structure(self):
        doc = docx.Document(self.doc_path)
        all_text_list = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    all_text_list.append(cell.text)
        full_text = "\n".join(all_text_list)
        
        # Verify university and author information
        self.assertIn("TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI", full_text)
        self.assertIn("Hoàng Quang Tùng", full_text)
        self.assertIn("2311060082", full_text)
        self.assertIn("SMART DESK ASSISTANT", full_text)
        
        # Verify key chapters
        self.assertIn("CHƯƠNG 1", full_text)
        self.assertIn("CHƯƠNG 2", full_text)
        self.assertIn("CHƯƠNG 3", full_text)
        self.assertIn("KẾT LUẬN VÀ KIẾN NGHỊ", full_text)
        self.assertIn("TÀI LIỆU THAM KHẢO", full_text)
        
        # Verify technical content
        self.assertIn("LD2410C", full_text)
        self.assertIn("ESP-NOW", full_text)
        self.assertIn("VBUS", full_text)
        
        # Verify tables count
        self.assertGreaterEqual(len(doc.tables), 6, "Report must contain at least 6 structured tables")

if __name__ == '__main__':
    unittest.main()
